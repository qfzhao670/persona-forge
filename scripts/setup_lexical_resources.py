#!/usr/bin/env python3
"""Download, verify, and prepare fixed lexical resources for W1/W2/G1."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import tempfile
import urllib.request
import zipfile
from pathlib import Path


ROOT = Path(__file__).resolve().parent.parent
DEFAULT_RESOURCE_DIR = ROOT / "resources"
FUNCTION_WORD_SOURCE = ROOT / "resources" / "function_words_v1.json"
LEXICON_VERSION = "subtlex_us_pos_zipf_2013+nltk_3.9.2+wordnet_3.0+surface_then_lemma_v1"
FUNCTION_WORD_VERSION = "reddit_g1_function_words_en_v1"
G1_CATEGORIES = (
    "article", "preposition", "personal_pronoun", "impersonal_pronoun",
    "auxiliary_verb", "conjunction", "adverb", "negation",
)
SOURCES = {
    "subtlexus1.zip": {
        "url": "https://www.ugent.be/plone_portal/pp/experimentele-psychologie/en/research/documents/subtlexus/subtlexus1.zip",
        "sha256": "458128f90a28c4f396cb2a5b23ac93c56f745ee8cfca9be2afedad4091d15090",
    },
    "averaged_perceptron_tagger_eng.zip": {
        "url": "https://raw.githubusercontent.com/nltk/nltk_data/gh-pages/packages/taggers/averaged_perceptron_tagger_eng.zip",
        "sha256": "6025f530624335c67d6547d44757b357b4e79bae030a0383e9887a92c1718f0b",
    },
    "wordnet.zip": {
        "url": "https://raw.githubusercontent.com/nltk/nltk_data/gh-pages/packages/corpora/wordnet.zip",
        "sha256": "cbda5ea6eef7f36a97a43d4a75f85e07fccbb4f23657d27b4ccbc93e2646ab59",
    },
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_tree(path: Path) -> str:
    digest = hashlib.sha256()
    for child in sorted(item for item in path.rglob("*") if item.is_file()):
        digest.update(child.relative_to(path).as_posix().encode("utf-8"))
        digest.update(b"\0")
        with child.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
    return digest.hexdigest()


def obtain_source(name: str, cache_dir: Path | None, work_dir: Path) -> Path:
    spec = SOURCES[name]
    cached = cache_dir / name if cache_dir is not None else None
    target = work_dir / name
    if cached is not None and cached.exists():
        shutil.copy2(cached, target)
    else:
        print(f"downloading {spec['url']}")
        with urllib.request.urlopen(spec["url"], timeout=120) as response:
            with target.open("wb") as output:
                shutil.copyfileobj(response, output)
    actual = sha256_file(target)
    if actual != spec["sha256"]:
        raise RuntimeError(f"{name} SHA-256 mismatch: {actual}")
    return target


def safe_extract(source: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    base = destination.resolve()
    with zipfile.ZipFile(source) as archive:
        for member in archive.infolist():
            target = (destination / member.filename).resolve()
            if target != base and base not in target.parents:
                raise RuntimeError(f"unsafe ZIP path: {member.filename}")
        archive.extractall(destination)


def build_subtlex(source_zip: Path, destination: Path) -> tuple[int, str]:
    try:
        import openpyxl
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl==3.1.5 is required; install requirements.txt") from exc
    if openpyxl.__version__ != "3.1.5":
        raise RuntimeError(f"openpyxl version must be 3.1.5, got {openpyxl.__version__}")

    with tempfile.TemporaryDirectory(prefix="subtlex-xlsx-") as directory:
        extracted = Path(directory)
        safe_extract(source_zip, extracted)
        candidates = list(extracted.glob("*.xlsx"))
        if len(candidates) != 1:
            raise RuntimeError("SUBTLEX-US ZIP must contain exactly one XLSX file")
        worksheet = load_workbook(candidates[0], read_only=True, data_only=True).active
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows)
        try:
            word_index = header.index("Word")
            zipf_index = header.index("Zipf-value")
        except ValueError as exc:
            raise RuntimeError("SUBTLEX-US columns Word/Zipf-value not found") from exc

        values: dict[str, float] = {}
        for row_number, row in enumerate(rows, 2):
            word = row[word_index]
            zipf = row[zipf_index]
            if not isinstance(word, str) or not isinstance(zipf, (int, float)):
                raise RuntimeError(f"invalid SUBTLEX-US row {row_number}")
            key = word.lower()
            if key in values:
                raise RuntimeError(f"duplicate lowercase SUBTLEX-US word: {key}")
            values[key] = float(zipf)

    if len(values) != 74286:
        raise RuntimeError(f"expected 74286 SUBTLEX-US rows, got {len(values)}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(("word", "zipf"))
        for word in sorted(values):
            writer.writerow((word, format(values[word], ".15g")))
    return len(values), sha256_file(destination)


def prepare_function_words(destination: Path) -> tuple[int, int, str]:
    try:
        value = json.loads(FUNCTION_WORD_SOURCE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"invalid tracked function-word lexicon: {exc}") from exc
    if value.get("format_version") != 1:
        raise RuntimeError("unsupported function-word lexicon format")
    if value.get("lexicon_version") != FUNCTION_WORD_VERSION:
        raise RuntimeError("function-word lexicon version mismatch")
    categories = value.get("categories")
    if not isinstance(categories, dict) or tuple(categories) != G1_CATEGORIES:
        raise RuntimeError("function-word categories or category order mismatch")
    memberships = 0
    unique_words: set[str] = set()
    for category in G1_CATEGORIES:
        words = categories[category]
        if (
            not isinstance(words, list)
            or not words
            or not all(isinstance(word, str) and word for word in words)
            or words != sorted(set(words))
        ):
            raise RuntimeError(f"invalid function-word category: {category}")
        memberships += len(words)
        unique_words.update(words)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.resolve() != FUNCTION_WORD_SOURCE.resolve():
        shutil.copy2(FUNCTION_WORD_SOURCE, destination)
    return memberships, len(unique_words), sha256_file(destination)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--resource-dir", type=Path, default=DEFAULT_RESOURCE_DIR)
    parser.add_argument(
        "--source-cache", type=Path, default=None,
        help="optional directory containing the three already-downloaded ZIP files",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    resource_dir = args.resource_dir.resolve()
    cache_dir = args.source_cache.resolve() if args.source_cache else None
    resource_dir.mkdir(parents=True, exist_ok=True)
    function_memberships, function_unique_words, function_hash = prepare_function_words(
        resource_dir / "function_words_v1.json"
    )

    with tempfile.TemporaryDirectory(prefix="reddit-lexical-resources-") as directory:
        work_dir = Path(directory)
        sources = {
            name: obtain_source(name, cache_dir, work_dir)
            for name in SOURCES
        }
        entry_count, normalized_hash = build_subtlex(
            sources["subtlexus1.zip"], resource_dir / "subtlex_us_zipf.tsv"
        )
        nltk_data = resource_dir / "nltk_data"
        tagger_dir = nltk_data / "taggers" / "averaged_perceptron_tagger_eng"
        wordnet_dir = nltk_data / "corpora" / "wordnet"
        if tagger_dir.exists():
            shutil.rmtree(tagger_dir)
        if wordnet_dir.exists():
            shutil.rmtree(wordnet_dir)
        safe_extract(sources["averaged_perceptron_tagger_eng.zip"], nltk_data / "taggers")
        safe_extract(sources["wordnet.zip"], nltk_data / "corpora")

    manifest = {
        "format_version": 2,
        "lexicon_version": LEXICON_VERSION,
        "subtlex_us": {
            "source_url": SOURCES["subtlexus1.zip"]["url"],
            "source_sha256": SOURCES["subtlexus1.zip"]["sha256"],
            "entry_count": entry_count,
            "normalized_file": "subtlex_us_zipf.tsv",
            "normalized_sha256": normalized_hash,
        },
        "nltk": {
            "package_version": "3.9.2",
            "tagger": {
                "name": "averaged_perceptron_tagger_eng",
                "source_url": SOURCES["averaged_perceptron_tagger_eng.zip"]["url"],
                "source_sha256": SOURCES["averaged_perceptron_tagger_eng.zip"]["sha256"],
                "installed_sha256": sha256_tree(
                    resource_dir / "nltk_data" / "taggers" / "averaged_perceptron_tagger_eng"
                ),
            },
            "lemmatizer": {
                "name": "wordnet_3.0",
                "source_url": SOURCES["wordnet.zip"]["url"],
                "source_sha256": SOURCES["wordnet.zip"]["sha256"],
                "installed_sha256": sha256_tree(
                    resource_dir / "nltk_data" / "corpora" / "wordnet"
                ),
            },
        },
        "function_words": {
            "file": "function_words_v1.json",
            "version": FUNCTION_WORD_VERSION,
            "sha256": function_hash,
            "memberships": function_memberships,
            "unique_words": function_unique_words,
            "liwc_compatible": False,
        },
        "lookup_policy": "surface_then_apostrophe_stripped_then_lemma_v1",
    }
    (resource_dir / "lexical_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        f"prepared {entry_count} SUBTLEX-US entries and "
        f"{function_unique_words} G1 function words in {resource_dir}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
